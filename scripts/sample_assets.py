from __future__ import annotations

import argparse
import shutil
import tempfile
import textwrap
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from pdf_translator.config import AppConfig
from pdf_translator.workflow import TranslationWorkflow


def create_sample_pdf(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    image_path = path.with_name("_sample_figure.png")
    image = Image.new("RGB", (900, 420), "#EFF6FF")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((30, 30, 870, 390), radius=28, outline="#2563EB", width=6)
    draw.line((90, 320, 260, 220, 430, 260, 620, 120, 800, 170), fill="#1D4ED8", width=10)
    draw.text((70, 55), "Sample figure: model accuracy", fill="#0F172A")
    image.save(image_path)

    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    pdf.setTitle("Local PDF Translator Sample")

    pdf.setFont("Helvetica-Bold", 22)
    pdf.drawString(22 * mm, height - 28 * mm, "A Practical Study of PDF Translation")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(22 * mm, height - 37 * mm, "Example Author - author@example.com")

    body = (
        "This study evaluates a local-first PDF translation workflow. "
        "The experiment uses 25 mg of material at 20 C and reports an accuracy of 96%. "
        "Project documentation is available at https://example.com/research. "
        "The simplified energy relation is E = mc^2 and the method follows reference [12]."
    )
    text = pdf.beginText(22 * mm, height - 52 * mm)
    text.setFont("Times-Roman", 11)
    text.setLeading(16)
    for line in textwrap.wrap(body, width=95):
        text.textLine(line)
    pdf.drawText(text)

    pdf.drawImage(
        str(image_path),
        28 * mm,
        height - 150 * mm,
        width=154 * mm,
        height=72 * mm,
        preserveAspectRatio=True,
        mask="auto",
    )
    pdf.setFont("Helvetica-Oblique", 9)
    pdf.drawCentredString(
        width / 2,
        height - 155 * mm,
        "Figure 1. Accuracy improves during training.",
    )

    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(22 * mm, height - 174 * mm, "1. Introduction")
    pdf.setFont("Times-Roman", 10.5)
    intro = (
        "PDF files combine text, fonts, vector graphics, raster images, and page geometry. "
        "A safe translator should preserve content even when the original layout cannot be reused. "
        "For that reason, this MVP generates readable flow-layout pages and a bilingual review file."
    )
    y = height - 184 * mm
    for line in textwrap.wrap(intro, width=104):
        pdf.drawString(22 * mm, y, line)
        y -= 5.2 * mm
    pdf.setFont("Helvetica", 8)
    pdf.drawCentredString(width / 2, 10 * mm, "1")
    pdf.showPage()

    pdf.setFont("Helvetica-Bold", 17)
    pdf.drawString(20 * mm, height - 24 * mm, "2. Two-column extraction example")
    left = (
        "The left column contains the first part of the reading order. "
        "Segments receive stable task-local identifiers and numeric row keys. "
        "URLs, e-mail addresses, references, code, and units are protected before export."
    )
    right = (
        "The right column contains the second part. Returned spreadsheets are validated for "
        "missing rows, duplicate keys, empty translations, changed numbers, and damaged placeholders. "
        "Blocking errors must be fixed before PDF generation."
    )
    for x, paragraph in ((20 * mm, left), (108 * mm, right)):
        text = pdf.beginText(x, height - 38 * mm)
        text.setFont("Times-Roman", 10)
        text.setLeading(14)
        for line in textwrap.wrap(paragraph, width=48):
            text.textLine(line)
        pdf.drawText(text)

    table_x = 22 * mm
    table_y = height - 125 * mm
    col_widths = [45 * mm, 45 * mm, 45 * mm]
    row_height = 10 * mm
    rows = [
        ["Mode", "Overlap risk", "Readable"],
        ["Fixed box", "High", "Sometimes"],
        ["Flow layout", "Low", "Yes"],
    ]
    pdf.setStrokeColor(colors.HexColor("#94A3B8"))
    for row_index, row in enumerate(rows):
        y_top = table_y - row_index * row_height
        x = table_x
        for col_index, value in enumerate(row):
            pdf.rect(x, y_top - row_height, col_widths[col_index], row_height)
            pdf.setFont("Helvetica-Bold" if row_index == 0 else "Helvetica", 9)
            pdf.drawString(x + 3 * mm, y_top - 6.5 * mm, value)
            x += col_widths[col_index]

    pdf.setFont("Helvetica", 9)
    pdf.drawString(22 * mm, height - 168 * mm, "Contact: qa@example.org")
    pdf.drawString(22 * mm, height - 174 * mm, "DOI: 10.1234/example.2026.001")
    pdf.drawCentredString(width / 2, 10 * mm, "2")
    pdf.save()
    image_path.unlink(missing_ok=True)
    return path


def mock_translate_xlsx(source: Path, destination: Path) -> Path:
    workbook = load_workbook(source)
    sheet = workbook["Translate"]
    sheet.title = "翻译结果"
    sheet["A1"] = "行键 - 请勿修改"
    sheet["B1"] = "需要翻译的文本"
    for row in range(2, sheet.max_row + 1):
        source_text = str(sheet.cell(row=row, column=2).value or "")
        sheet.cell(row=row, column=2).value = f"模拟中文译文：{source_text}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def build_assets(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    source_pdf = create_sample_pdf(output_dir / "sample_source.pdf")
    with tempfile.TemporaryDirectory(prefix="pdf-translator-sample-") as temporary:
        workflow = TranslationWorkflow(AppConfig.from_env(Path(temporary) / "data"))
        task = workflow.create_task(source_pdf)
        _, packages = workflow.export_packages(task.task_id, Path(temporary) / "exports")
        shutil.copy2(packages[0], output_dir / "sample_translation_package.xlsx")
        mock_translate_xlsx(
            packages[0],
            output_dir / "sample_translated_package.xlsx",
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="examples")
    args = parser.parse_args()
    build_assets(Path(args.output).resolve())


if __name__ == "__main__":
    main()
