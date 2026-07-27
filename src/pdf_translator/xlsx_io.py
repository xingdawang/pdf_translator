from __future__ import annotations

import re
import tempfile
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from .exceptions import PackageError
from .models import ExportPackage, TranslationTask, ValidationIssue
from .utils import safe_stem, utc_now


@dataclass
class ImportResult:
    filename: str
    package_id: str | None
    imported_segments: int
    exact_matches: int
    positional_matches: int
    issues: list[ValidationIssue] = field(default_factory=list)


class XLSXExporter:
    def __init__(
        self,
        max_rows: int = 1_000_000,
        max_characters: int = 100_000_000,
        max_file_bytes: int = 9_500_000,
    ):
        self.max_rows = max_rows
        self.max_characters = max_characters
        self.max_file_bytes = max_file_bytes

    def export(self, task: TranslationTask, output_dir: Path) -> list[ExportPackage]:
        output_dir.mkdir(parents=True, exist_ok=True)
        segments = task.translatable_segments
        if not segments:
            raise PackageError("当前任务没有需要翻译的段落。")

        chunks: list[list[Any]] = []
        current: list[Any] = []
        current_characters = 0
        for segment in segments:
            length = len(segment.protected_text)
            if current and (
                len(current) >= self.max_rows
                or current_characters + length > self.max_characters
            ):
                chunks.append(current)
                current = []
                current_characters = 0
            current.append(segment)
            current_characters += length
        if current:
            chunks.append(current)

        chunks = self._split_oversized_chunks(task, chunks, output_dir)
        stem = safe_stem(task.source_filename)
        packages: list[ExportPackage] = []
        total = len(chunks)
        for package_index, chunk in enumerate(chunks, start=1):
            package_id = (
                f"PKG-{task.task_id}-{package_index:03d}-{uuid.uuid4().hex[:8].upper()}"
            )
            filename = (
                f"{stem}_{task.task_id}_google_translate_"
                f"{package_index:03d}_of_{total:03d}.xlsx"
            )
            path = output_dir / filename
            self._write_workbook(
                task=task,
                package_id=package_id,
                package_index=package_index,
                package_total=total,
                segments=chunk,
                path=path,
            )
            packages.append(
                ExportPackage(
                    package_id=package_id,
                    filename=filename,
                    package_index=package_index,
                    package_total=total,
                    row_keys=[segment.row_key for segment in chunk],
                    created_at=utc_now(),
                )
            )
        return packages

    def _split_oversized_chunks(
        self,
        task: TranslationTask,
        chunks: list[list[Any]],
        output_dir: Path,
    ) -> list[list[Any]]:
        """Keep each actual XLSX below Google's document-upload limit."""
        if self.max_file_bytes <= 0:
            return chunks
        refined: list[list[Any]] = []
        with tempfile.TemporaryDirectory(
            prefix=".xlsx-size-probe-",
            dir=output_dir,
        ) as temporary_dir:
            probe_dir = Path(temporary_dir)
            for chunk in chunks:
                remaining = chunk
                while remaining:
                    full_size = self._measure_workbook(
                        task,
                        remaining,
                        probe_dir,
                    )
                    if full_size <= self.max_file_bytes:
                        refined.append(remaining)
                        break
                    if len(remaining) == 1:
                        raise PackageError(
                            "单个翻译段落生成的 XLSX 已超过 9.5 MB，"
                            "请缩短该段原文后重试。"
                        )
                    split_at = self._largest_fitting_prefix(
                        task,
                        remaining,
                        probe_dir,
                    )
                    refined.append(remaining[:split_at])
                    remaining = remaining[split_at:]
        return refined

    def _largest_fitting_prefix(
        self,
        task: TranslationTask,
        segments: list[Any],
        probe_dir: Path,
    ) -> int:
        low = 1
        high = len(segments) - 1
        best = 0
        while low <= high:
            middle = (low + high) // 2
            size = self._measure_workbook(
                task,
                segments[:middle],
                probe_dir,
            )
            if size <= self.max_file_bytes:
                best = middle
                low = middle + 1
            else:
                high = middle - 1
        if best:
            return best
        raise PackageError(
            "单个翻译段落生成的 XLSX 已超过 9.5 MB，"
            "请缩短该段原文后重试。"
        )

    def _measure_workbook(
        self,
        task: TranslationTask,
        segments: list[Any],
        probe_dir: Path,
    ) -> int:
        path = probe_dir / "probe.xlsx"
        self._write_workbook(
            task=task,
            package_id=f"PKG-{task.task_id}-SIZE-PROBE",
            package_index=1,
            package_total=1,
            segments=segments,
            path=path,
        )
        return path.stat().st_size

    @staticmethod
    def _write_workbook(
        task: TranslationTask,
        package_id: str,
        package_index: int,
        package_total: int,
        segments: list[Any],
        path: Path,
    ) -> None:
        workbook = Workbook(write_only=False)
        sheet = workbook.active
        sheet.title = "Translate"
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:B{len(segments) + 1}"

        header_fill = PatternFill("solid", fgColor="1D4ED8")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        light_border = Border(
            bottom=Side(style="thin", color="CBD5E1")
        )

        sheet.append(["Row Key - do not edit", "Text to translate"])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 26

        for row_number, segment in enumerate(segments, start=2):
            sheet.append([segment.row_key, segment.protected_text])
            key_cell = sheet.cell(row=row_number, column=1)
            text_cell = sheet.cell(row=row_number, column=2)
            key_cell.number_format = "@"
            key_cell.protection = Protection(locked=True)
            key_cell.alignment = Alignment(vertical="top", horizontal="center")
            text_cell.alignment = Alignment(
                wrap_text=True, vertical="top", horizontal="left"
            )
            key_cell.border = light_border
            text_cell.border = light_border
            sheet.row_dimensions[row_number].height = min(
                120, max(22, 18 * (1 + len(segment.protected_text) // 100))
            )

        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 105
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"

        metadata = workbook.create_sheet("Package Info")
        metadata.sheet_state = "hidden"
        metadata_rows = [
            ("format", "local-pdf-translator-xlsx"),
            ("format_version", "2"),
            (
                "document_ir_version",
                task.document_ir.version if task.document_ir else 0,
            ),
            ("task_id", task.task_id),
            ("package_id", package_id),
            ("package_index", package_index),
            ("package_total", package_total),
            ("source_filename", task.source_filename),
            ("source_file_hash", task.source_file_hash),
            ("segment_count", len(segments)),
            ("source_language", "en"),
            ("target_language", "zh-CN"),
            ("created_at", utc_now()),
        ]
        for row in metadata_rows:
            metadata.append(row)
        metadata.column_dimensions["A"].width = 24
        metadata.column_dimensions["B"].width = 80

        workbook.properties.title = f"PDF translation package {package_index}/{package_total}"
        workbook.properties.subject = package_id
        workbook.properties.creator = "Local PDF Translator"

        temporary = path.with_suffix(".xlsx.tmp")
        workbook.save(temporary)
        temporary.replace(path)


class XLSXImporter:
    def import_file(
        self, task: TranslationTask, path: Path
    ) -> ImportResult:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise PackageError(f"无法读取 XLSX 文件 {path.name}：{exc}") from exc

        try:
            package_id = self._find_package_id(workbook)
            package = self._match_package(task, package_id, path.name)
            expected_keys = package.row_keys if package else [
                segment.row_key for segment in task.translatable_segments
            ]
            expected_set = set(expected_keys)

            sheet, rows, exact_key_count = self._find_translation_rows(
                workbook, expected_set
            )
            if sheet is None or (exact_key_count <= 0 and not package):
                raise PackageError(
                    f"{path.name} 中找不到翻译数据表，或者 Row Key 已全部丢失。"
                )

            translations: dict[str, str] = {}
            duplicates: list[str] = []
            exact_matches = 0
            positional_matches = 0

            for row in rows:
                key_index = self._find_key_index(row, expected_set)
                if key_index is None:
                    continue
                key = self._normalize_row_key(row[key_index])
                if key is None:
                    continue
                if key in translations:
                    duplicates.append(key)
                    continue
                translations[key] = self._translation_from_row(row, key_index)
                exact_matches += 1

            issues: list[ValidationIssue] = []
            for key in sorted(set(duplicates)):
                issues.append(
                    ValidationIssue(
                        severity="blocking",
                        code="DUPLICATE_ROW_KEY",
                        message=f"翻译文件中 Row Key {key} 出现多次。",
                    )
                )

            missing_keys = [key for key in expected_keys if key not in translations]
            if missing_keys and package_id and package:
                candidate_rows = [
                    row
                    for row in rows
                    if any(value not in (None, "") for value in row)
                ]
                if len(candidate_rows) == len(expected_keys) and exact_key_count < len(expected_keys) / 2:
                    translations.clear()
                    for key, row in zip(expected_keys, candidate_rows, strict=True):
                        translations[key] = self._translation_from_positional_row(row)
                    exact_matches = 0
                    positional_matches = len(expected_keys)
                    missing_keys = []
                    issues.append(
                        ValidationIssue(
                            severity="warning",
                            code="POSITIONAL_MATCH",
                            message=(
                                "Row Key 大部分被修改；已依据匹配的 Package ID 和原始行顺序恢复。"
                                "请重点抽查译文对应关系。"
                            ),
                        )
                    )

            segment_by_key = {
                segment.row_key: segment for segment in task.translatable_segments
            }
            imported = 0
            for key, translation in translations.items():
                segment = segment_by_key.get(key)
                if segment is None:
                    continue
                segment.translated_text = translation.strip()
                segment.status = "translated" if translation.strip() else "empty"
                segment.fallback_reason = None
                imported += 1

            for key in missing_keys[:100]:
                segment = segment_by_key.get(key)
                issues.append(
                    ValidationIssue(
                        severity="blocking",
                        code="MISSING_ROW_KEY",
                        message=f"翻译文件缺少 Row Key {key}。",
                        segment_id=segment.segment_id if segment else None,
                        page_number=segment.page_number if segment else None,
                    )
                )
            if len(missing_keys) > 100:
                issues.append(
                    ValidationIssue(
                        severity="blocking",
                        code="MISSING_ROW_KEY_SUMMARY",
                        message=f"另有 {len(missing_keys) - 100} 个缺失 Row Key 未逐条列出。",
                    )
                )

            return ImportResult(
                filename=path.name,
                package_id=package_id,
                imported_segments=imported,
                exact_matches=exact_matches,
                positional_matches=positional_matches,
                issues=issues,
            )
        finally:
            workbook.close()

    @staticmethod
    def _find_package_id(workbook: Any) -> str | None:
        subject = getattr(workbook.properties, "subject", None)
        if subject and str(subject).startswith("PKG-"):
            return str(subject)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row or 1, 30),
                min_col=1,
                max_col=min(sheet.max_column or 1, 8),
                values_only=True,
            ):
                for value in row:
                    if value and str(value).startswith("PKG-"):
                        return str(value).strip()
        return None

    @staticmethod
    def _match_package(
        task: TranslationTask, package_id: str | None, filename: str
    ) -> ExportPackage | None:
        if package_id:
            for package in task.export_packages:
                if package.package_id == package_id:
                    return package
            raise PackageError(
                f"{filename} 不属于任务 {task.task_id}，Package ID 不匹配。"
            )

        filename_match = re.search(r"_(\d{3})_of_(\d{3})\.xlsx$", filename)
        if filename_match:
            package_index = int(filename_match.group(1))
            candidates = [
                package
                for package in task.export_packages
                if package.package_index == package_index
            ]
            if len(candidates) == 1:
                return candidates[0]
        return None

    def _find_translation_rows(
        self, workbook: Any, expected_keys: set[str]
    ) -> tuple[Any | None, list[tuple[Any, ...]], int]:
        best_sheet = None
        best_rows: list[tuple[Any, ...]] = []
        best_score = (-1, -1, -1, -1)
        best_count = -1
        for sheet in workbook.worksheets:
            rows = list(
                sheet.iter_rows(
                    min_row=2,
                    max_row=sheet.max_row,
                    min_col=1,
                    max_col=min(max(sheet.max_column or 2, 2), 8),
                    values_only=True,
                )
            )
            matching_keys = [
                self._normalize_row_key(value)
                for row in rows
                for value in row
                if self._normalize_row_key(value) in expected_keys
            ]
            count = sum(
                self._find_key_index(row, expected_keys) is not None for row in rows
            )
            unique_count = len(set(matching_keys))
            nonempty_rows = sum(
                any(value not in (None, "") for value in row) for row in rows
            )
            visible = int(getattr(sheet, "sheet_state", "visible") == "visible")
            # The translated data sheet is expected to stay visible. Metadata may
            # contain small integers that resemble Row Keys, so visibility and
            # distinct-key coverage take precedence over raw match count.
            score = (visible, unique_count, count, nonempty_rows)
            if score > best_score:
                best_sheet = sheet
                best_rows = rows
                best_count = count
                best_score = score
        if best_sheet is None or not best_rows:
            return None, [], 0
        return best_sheet, best_rows, best_count

    @staticmethod
    def _normalize_row_key(value: Any) -> str | None:
        if isinstance(value, bool) or value is None:
            return None
        if isinstance(value, int):
            return f"{value:08d}"
        if isinstance(value, float) and value.is_integer():
            return f"{int(value):08d}"
        text = str(value).strip()
        if text.isdigit():
            return text.zfill(8)
        paragraph_match = re.fullmatch(
            r"P(\d{1,8})-S(\d{1,10})",
            text,
            re.IGNORECASE,
        )
        if paragraph_match:
            return (
                f"P{int(paragraph_match.group(1)):04d}-"
                f"S{int(paragraph_match.group(2)):06d}"
            )
        match = re.fullmatch(r"R?(\d{1,8})", text, re.IGNORECASE)
        return match.group(1).zfill(8) if match else None

    def _find_key_index(
        self, row: tuple[Any, ...], expected_keys: set[str]
    ) -> int | None:
        for index, value in enumerate(row):
            key = self._normalize_row_key(value)
            if key in expected_keys:
                return index
        return None

    @staticmethod
    def _translation_from_row(row: tuple[Any, ...], key_index: int) -> str:
        for value in row[key_index + 1 :]:
            if value is not None:
                return str(value)
        return ""

    @staticmethod
    def _translation_from_positional_row(row: tuple[Any, ...]) -> str:
        values = [value for value in row if value not in (None, "")]
        if not values:
            return ""
        if len(values) == 1:
            return str(values[0])
        return str(values[-1])
