from pathlib import Path

import pytest
from openpyxl import load_workbook

from pdf_translator.config import AppConfig
from pdf_translator.exceptions import PackageError
from pdf_translator.models import (
    Segment,
    TaskSettings,
    TranslationTask,
)
from pdf_translator.placeholders import PlaceholderService
from pdf_translator.repository import TaskRepository
from pdf_translator.utils import sha256_text, utc_now
from pdf_translator.validation import TranslationValidator
from pdf_translator.workflow import TranslationWorkflow
from pdf_translator.xlsx_io import XLSXExporter, XLSXImporter


def make_task(task_id: str = "20260726-testtask") -> TranslationTask:
    service = PlaceholderService()
    sources = [
        "Introduction",
        "Visit https://example.com and use 25 mg.",
        "Contact qa@example.org and see reference [12].",
    ]
    segments = []
    for index, source in enumerate(sources, start=1):
        segment_id = f"P0001-S{index:06d}"
        protected = service.protect(source, task_id, segment_id)
        segments.append(
            Segment(
                segment_id=segment_id,
                row_key=f"{index:08d}",
                page_number=1,
                reading_order=index,
                block_type="paragraph",
                source_text=source,
                protected_text=protected.protected_text,
                source_text_hash=sha256_text(source),
                bbox=[0, 0, 100, 20],
                placeholders=protected.placeholders,
            )
        )
    now = utc_now()
    return TranslationTask(
        task_id=task_id,
        source_path="/tmp/source.pdf",
        source_filename="source.pdf",
        source_file_hash="a" * 64,
        source_size_bytes=100,
        created_at=now,
        updated_at=now,
        status="analyzed",
        page_count=1,
        text_page_count=1,
        scanned_page_count=0,
        image_count=0,
        parser_version="test",
        settings=TaskSettings(),
        segments=segments,
    )


def translate_workbook(source: Path, destination: Path) -> None:
    workbook = load_workbook(source)
    sheet = workbook["Translate"]
    sheet.title = "翻译结果"
    sheet["A1"] = "行键"
    sheet["B1"] = "译文"
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=2).value = (
            "中文译文：" + str(sheet.cell(row=row, column=2).value)
        )
    workbook.save(destination)


def test_export_import_survives_translated_headers_and_sheet_name(tmp_path):
    task = make_task()
    packages = XLSXExporter().export(task, tmp_path)
    task.export_packages = packages
    source = tmp_path / packages[0].filename
    translated = tmp_path / "returned.xlsx"
    translate_workbook(source, translated)

    result = XLSXImporter().import_file(task, translated)
    report = TranslationValidator().validate(task, result.issues)
    assert result.exact_matches == 3
    assert result.imported_segments == 3
    assert report.can_generate


def test_missing_row_is_detected(tmp_path):
    task = make_task()
    packages = XLSXExporter().export(task, tmp_path)
    task.export_packages = packages
    source = tmp_path / packages[0].filename
    translated = tmp_path / "missing.xlsx"
    translate_workbook(source, translated)
    workbook = load_workbook(translated)
    workbook["翻译结果"].delete_rows(3)
    workbook.save(translated)

    result = XLSXImporter().import_file(task, translated)
    assert any(issue.code == "MISSING_ROW_KEY" for issue in result.issues)


def test_workflow_missing_translation_falls_back_to_source_and_is_ready(tmp_path):
    workflow = TranslationWorkflow(AppConfig.from_env(tmp_path / "data"))
    task = make_task()
    workflow.repository.save(task)
    task, packages = workflow.export_packages(task.task_id)
    translated = tmp_path / "missing-returned.xlsx"
    translate_workbook(packages[0], translated)
    workbook = load_workbook(translated)
    workbook["翻译结果"].delete_rows(3)
    workbook.save(translated)

    repaired, _, report = workflow.import_packages(
        task.task_id,
        [translated],
    )

    fallback = repaired.segments[1]
    assert report.can_generate
    assert fallback.status == "source_fallback"
    assert fallback.translated_text == fallback.source_text
    assert fallback.fallback_reason == "翻译件中缺少对应译文"
    assert any(
        issue.code == "SOURCE_FALLBACK"
        and issue.segment_id == fallback.segment_id
        for issue in report.issues
    )


def test_workflow_broken_critical_placeholder_falls_back_to_source(tmp_path):
    workflow = TranslationWorkflow(AppConfig.from_env(tmp_path / "data"))
    task = make_task()
    workflow.repository.save(task)
    task, packages = workflow.export_packages(task.task_id)
    translated = tmp_path / "broken-placeholder.xlsx"
    translate_workbook(packages[0], translated)
    workbook = load_workbook(translated)
    workbook["翻译结果"].cell(row=3, column=2).value = "链接和占位符均已丢失"
    workbook.save(translated)

    repaired, _, report = workflow.import_packages(
        task.task_id,
        [translated],
    )

    fallback = repaired.segments[1]
    assert report.can_generate
    assert fallback.status == "source_fallback"
    assert fallback.translated_text == fallback.source_text
    assert "关键内容无法安全恢复" in fallback.fallback_reason


def test_duplicate_row_key_is_detected(tmp_path):
    task = make_task()
    packages = XLSXExporter().export(task, tmp_path)
    task.export_packages = packages
    source = tmp_path / packages[0].filename
    translated = tmp_path / "duplicate.xlsx"
    translate_workbook(source, translated)
    workbook = load_workbook(translated)
    sheet = workbook["翻译结果"]
    sheet.cell(row=3, column=1).value = sheet.cell(row=2, column=1).value
    workbook.save(translated)

    result = XLSXImporter().import_file(task, translated)
    assert any(issue.code == "DUPLICATE_ROW_KEY" for issue in result.issues)


def test_position_fallback_requires_matching_package_id_and_row_count(tmp_path):
    task = make_task()
    packages = XLSXExporter().export(task, tmp_path)
    task.export_packages = packages
    source = tmp_path / packages[0].filename
    translated = tmp_path / "position.xlsx"
    translate_workbook(source, translated)
    workbook = load_workbook(translated)
    sheet = workbook["翻译结果"]
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=1).value = f"changed-{row}"
    workbook.save(translated)

    result = XLSXImporter().import_file(task, translated)
    assert result.positional_matches == 3
    assert any(issue.code == "POSITIONAL_MATCH" for issue in result.issues)


def test_export_splits_large_task_into_numbered_packages(tmp_path):
    task = make_task()
    packages = XLSXExporter(max_rows=2).export(task, tmp_path)
    assert len(packages) == 2
    assert [package.package_total for package in packages] == [2, 2]
    assert all(task.task_id in package.filename for package in packages)
    assert packages[0].row_keys == ["00000001", "00000002"]
    assert packages[1].row_keys == ["00000003"]


def test_export_uses_actual_xlsx_size_and_avoids_legacy_row_splitting(tmp_path):
    task = make_task()
    template = task.segments[-1]
    for index in range(4, 3102):
        source = f"Unique translation row {index}: {sha256_text(str(index))}"
        task.segments.append(
            Segment(
                segment_id=f"P0001-S{index:06d}",
                row_key=f"{index:08d}",
                page_number=1,
                reading_order=index,
                block_type="paragraph",
                source_text=source,
                protected_text=source,
                source_text_hash=sha256_text(source),
                bbox=template.bbox,
            )
        )

    packages = XLSXExporter().export(task, tmp_path)

    assert len(packages) == 1
    assert (tmp_path / packages[0].filename).stat().st_size < 9_500_000


def test_wrong_task_package_is_rejected(tmp_path):
    first = make_task("20260726-first")
    packages = XLSXExporter().export(first, tmp_path)
    first.export_packages = packages
    translated = tmp_path / "wrong.xlsx"
    translate_workbook(tmp_path / packages[0].filename, translated)
    second = make_task("20260726-second")

    with pytest.raises(PackageError, match="不属于任务"):
        XLSXImporter().import_file(second, translated)


def test_empty_translation_is_blocking(tmp_path):
    task = make_task()
    task.segments[0].translated_text = ""
    for segment in task.segments[1:]:
        segment.translated_text = "中文译文：" + segment.protected_text
    report = TranslationValidator().validate(task)
    assert not report.can_generate
    assert any(issue.code == "EMPTY_TRANSLATION" for issue in report.issues)


def test_number_changes_are_ignored_by_default_and_can_be_reenabled():
    task = make_task()
    for segment in task.translatable_segments:
        segment.translated_text = "中文译文：" + segment.protected_text
    first = task.translatable_segments[0]
    first.source_text = "Version 2"
    first.translated_text = "版本 3"

    default_report = TranslationValidator().validate(task)

    assert not any(
        issue.code == "NUMBER_CHANGED" for issue in default_report.issues
    )

    task.settings.ignore_number_warnings = False
    strict_report = TranslationValidator().validate(task)

    assert any(issue.code == "NUMBER_CHANGED" for issue in strict_report.issues)


def test_missing_number_unit_placeholder_is_allowed_by_default():
    task = make_task()
    for segment in task.translatable_segments:
        segment.translated_text = "中文译文：" + segment.protected_text
    number_segment = task.translatable_segments[1]
    token = next(
        placeholder.token
        for placeholder in number_segment.placeholders
        if placeholder.kind == "number_unit"
    )
    number_segment.translated_text = number_segment.translated_text.replace(
        token,
        "25 毫克",
    )

    report = TranslationValidator().validate(task)

    assert report.can_generate
    assert any(
        issue.code == "PLACEHOLDER_CHANGED_ALLOWED"
        and issue.segment_id == number_segment.segment_id
        for issue in report.issues
    )


def test_review_confirmation_is_recorded_and_invalidated_by_edit(tmp_path):
    config = AppConfig.from_env(tmp_path / "data")
    workflow = TranslationWorkflow(config)
    task = make_task()
    for segment in task.translatable_segments:
        segment.translated_text = "中文译文：" + segment.protected_text
    workflow.repository.save(task)

    confirmed, report = workflow.confirm_review(task.task_id)

    assert report.can_generate
    assert confirmed.review_confirmation is not None
    assert confirmed.review_confirmation["translated_segments"] == 3

    confirmed.outputs = {"layout_pdf": "/tmp/stale-layout.pdf"}
    confirmed.last_generation = {"layout_cache_hit": True}
    workflow.repository.save(confirmed)
    first = confirmed.translatable_segments[0]
    edited, _ = workflow.update_translations(
        task.task_id,
        {first.segment_id: "修订译文：" + first.protected_text},
    )

    assert edited.review_confirmation is None
    assert edited.outputs == {}
    assert edited.last_generation is None
